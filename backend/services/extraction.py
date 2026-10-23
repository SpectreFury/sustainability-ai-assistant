import hashlib
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePath

import httpx
from langchain_core.messages import HumanMessage, SystemMessage
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from pypdf import PdfReader

from prompts.pdf_extraction import PDF_SYSTEM_PROMPT, PDF_USER_PROMPT
from prompts.spreadsheet_mapping import (
    SPREADSHEET_MAPPING_SYSTEM_PROMPT,
    spreadsheet_mapping_input,
)
from utils.cloudinary import create_signed_raw_download_url


class MetricColumnMapping(BaseModel):
    column_index: int = Field(ge=0, description="Zero-based spreadsheet column index")
    metric_name: str = Field(description="Canonical snake_case metric name")
    unit: str | None = Field(default=None, description="Unit specified in the header")


class SpreadsheetHeaderMapping(BaseModel):
    metric_columns: list[MetricColumnMapping] = Field(default_factory=list)
    metric_name_column: int | None = None
    metric_value_column: int | None = None
    period_start_column: int | None = None
    period_end_column: int | None = None
    financial_year_column: int | None = None
    unit_column: int | None = None


class PdfMetric(BaseModel):
    metric_name: str = Field(description="Metric name in snake_case")
    metric_value: str = Field(
        description="The exact printed number, as a string; do not calculate or round"
    )
    unit: str | None = None
    billing_period_start: str | None = None
    billing_period_end: str | None = None
    financial_year: str | None = None
    page_number: int = Field(ge=1, description="1-based page containing this value")


class PdfExtraction(BaseModel):
    metrics: list[PdfMetric] = Field(default_factory=list)


@dataclass(frozen=True)
class ExtractedMetric:
    metric_name: str
    metric_value: Decimal
    unit: str | None
    billing_period_start: str | None
    billing_period_end: str | None
    financial_year: str | None
    source_reference: str


_HEADER_SEARCH_ROWS = 30
_MAX_DOCUMENT_BYTES = int(os.getenv("EXTRACTION_MAX_BYTES", str(20 * 1024 * 1024)))
_UNIT_PATTERN = re.compile(
    r"\b(kwh|mwh|wh|kw|mw|inr|rupees?|lit(?:er|re)s?|kg|tonnes?|tons?)\b",
    re.IGNORECASE,
)


async def download_document(
    public_id: str,
    filename: str,
    expected_hash: str | None = None,
) -> bytes:
    file_format = PurePath(filename).suffix.lower().removeprefix(".")
    if file_format not in {"pdf", "xlsx"}:
        raise ValueError("Unsupported file type. Only .xlsx and .pdf are supported.")

    url = create_signed_raw_download_url(public_id, file_format)

    chunks: list[bytes] = []
    total_bytes = 0
    content_hash = hashlib.sha256()
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            async with client.stream("GET", url) as response:
                if not response.is_success:
                    cloudinary_error = response.headers.get("x-cld-error")
                    detail = f": {cloudinary_error}" if cloudinary_error else ""
                    raise ValueError(
                        "Cloudinary download failed with HTTP "
                        f"{response.status_code}{detail}."
                    )
                content_length = response.headers.get("content-length")
                if content_length and int(content_length) > _MAX_DOCUMENT_BYTES:
                    raise ValueError(
                        "Document exceeds the configured extraction size limit."
                    )

                async for chunk in response.aiter_bytes():
                    total_bytes += len(chunk)
                    if total_bytes > _MAX_DOCUMENT_BYTES:
                        raise ValueError(
                            "Document exceeds the configured extraction size limit."
                        )
                    chunks.append(chunk)
                    content_hash.update(chunk)
    except httpx.HTTPError:
        # The signed download URL contains a temporary API signature; don't leak it
        # through the worker's exception logging.
        raise ValueError("Cloudinary document download request failed.") from None

    if not chunks:
        raise ValueError("Downloaded document is empty.")
    if expected_hash and content_hash.hexdigest() != expected_hash:
        raise ValueError(
            "Downloaded document checksum does not match the uploaded file."
        )
    return b"".join(chunks)


async def extract_document(
    filename: str, content: bytes
) -> tuple[str, list[ExtractedMetric]]:
    extension = PurePath(filename).suffix.lower()
    if extension == ".xlsx":
        return "spreadsheet", await extract_spreadsheet(content)
    if extension == ".pdf":
        return "pdf_bill", await extract_pdf(content)
    raise ValueError("Unsupported document type. Only .xlsx and .pdf are supported.")


async def extract_spreadsheet(content: bytes) -> list[ExtractedMetric]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    records: list[ExtractedMetric] = []
    mapping_errors: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            header_row, headers = _find_header_row(worksheet)
            if header_row is None:
                continue

            mapping = _known_header_mapping(headers)
            if not _mapping_is_usable(mapping):
                try:
                    mapping = await _map_headers_with_groq(headers)
                except Exception as error:
                    mapping_errors.append(str(error))
                    continue

            try:
                _validate_mapping(mapping, len(headers))
            except ValueError as error:
                mapping_errors.append(str(error))
                continue

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=worksheet.max_row,
                    max_col=len(headers),
                    values_only=True,
                ),
                start=header_row + 1,
            ):
                if not any(
                    value is not None and str(value).strip() for value in row
                ):
                    continue

                source_reference = _row_reference(worksheet.title, row, row_number)
                period_start = _cell_text(
                    _get_column_value(row, mapping.period_start_column)
                )
                period_end = _cell_text(
                    _get_column_value(row, mapping.period_end_column)
                )
                financial_year = _financial_year(
                    _cell_text(_get_column_value(row, mapping.financial_year_column)),
                    period_start,
                )
                row_unit = _cell_text(_get_column_value(row, mapping.unit_column))

                if mapping.metric_columns:
                    for metric_column in mapping.metric_columns:
                        value = _to_decimal(row[metric_column.column_index])
                        if value is None:
                            continue
                        records.append(
                            ExtractedMetric(
                                metric_name=_metric_name(metric_column.metric_name),
                                metric_value=value,
                                unit=row_unit or metric_column.unit,
                                billing_period_start=period_start,
                                billing_period_end=period_end,
                                financial_year=financial_year,
                                source_reference=source_reference,
                            )
                        )
                elif (
                    mapping.metric_name_column is not None
                    and mapping.metric_value_column is not None
                ):
                    metric_name = _cell_text(row[mapping.metric_name_column])
                    value = _to_decimal(row[mapping.metric_value_column])
                    if metric_name and value is not None:
                        records.append(
                            ExtractedMetric(
                                metric_name=_metric_name(metric_name),
                                metric_value=value,
                                unit=row_unit,
                                billing_period_start=period_start,
                                billing_period_end=period_end,
                                financial_year=financial_year,
                                source_reference=source_reference,
                            )
                        )
    finally:
        workbook.close()

    if not records:
        if mapping_errors:
            raise ValueError(mapping_errors[0])
        raise ValueError("No numeric metrics were found in the spreadsheet.")
    return records


async def extract_pdf(content: bytes) -> list[ExtractedMetric]:
    reader = PdfReader(BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported.")
    page_count = len(reader.pages)
    if page_count == 0:
        raise ValueError("The PDF contains no pages.")

    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise ValueError("Set GEMINI_API_KEY to enable PDF extraction.")

    from langchain_google_genai import ChatGoogleGenerativeAI

    model = ChatGoogleGenerativeAI(
        model=os.getenv("GEMINI_MODEL", "gemini-3.8-flash"),
        api_key=api_key,
        temperature=0,
    ).with_structured_output(PdfExtraction)

    messages = [
        SystemMessage(content=PDF_SYSTEM_PROMPT),
        HumanMessage(
            content=[
                {"type": "text", "text": PDF_USER_PROMPT},
                {"type": "media", "mime_type": "application/pdf", "data": content},
            ]
        ),
    ]
    result = await model.ainvoke(messages)
    extraction = PdfExtraction.model_validate(result)

    records: list[ExtractedMetric] = []
    for metric in extraction.metrics:
        if metric.page_number > page_count:
            raise ValueError(
                f"Gemini returned page {metric.page_number}, but this PDF has "
                f"{page_count} pages."
            )
        value = _to_decimal(metric.metric_value)
        if value is None:
            raise ValueError("Gemini returned a metric without a valid numeric value.")
        period_start = _cell_text(metric.billing_period_start)
        records.append(
            ExtractedMetric(
                metric_name=_metric_name(metric.metric_name),
                metric_value=value,
                unit=_cell_text(metric.unit),
                billing_period_start=period_start,
                billing_period_end=_cell_text(metric.billing_period_end),
                financial_year=_financial_year(metric.financial_year, period_start),
                source_reference=f"Page {metric.page_number}",
            )
        )
    if not records:
        raise ValueError("No numeric metrics were found in the PDF.")
    return records


def _find_header_row(worksheet) -> tuple[int | None, list[str]]:
    max_column = worksheet.max_column or 1
    best_row: int | None = None
    best_headers: list[str] = []
    best_score = -1

    for row_index in range(1, min(worksheet.max_row or 1, _HEADER_SEARCH_ROWS) + 1):
        cells = next(
            worksheet.iter_rows(
                min_row=row_index,
                max_row=row_index,
                max_col=max_column,
                values_only=True,
            ),
            (),
        )
        non_empty = [
            value for value in cells if value is not None and str(value).strip()
        ]
        if len(non_empty) < 2:
            continue

        headers = [str(value).strip() if value is not None else "" for value in cells]
        recognized = sum(_is_known_header(header) for header in headers)
        score = recognized * 100 + len(non_empty)
        if score > best_score:
            best_row = row_index
            best_headers = headers
            best_score = score

    return best_row, best_headers


def _known_header_mapping(headers: list[str]) -> SpreadsheetHeaderMapping:
    mapping = SpreadsheetHeaderMapping()
    for column_index, header in enumerate(headers):
        normalized = _normalize_header(header)
        if not normalized:
            continue

        role = _header_role(normalized)
        if role == "period_start" and mapping.period_start_column is None:
            mapping.period_start_column = column_index
        elif role == "period_end" and mapping.period_end_column is None:
            mapping.period_end_column = column_index
        elif role == "financial_year" and mapping.financial_year_column is None:
            mapping.financial_year_column = column_index
        elif role == "unit" and mapping.unit_column is None:
            mapping.unit_column = column_index
        elif role == "metric_name" and mapping.metric_name_column is None:
            mapping.metric_name_column = column_index
        elif role == "metric_value" and mapping.metric_value_column is None:
            mapping.metric_value_column = column_index
        else:
            metric_name = _known_metric_name(normalized)
            if metric_name:
                mapping.metric_columns.append(
                    MetricColumnMapping(
                        column_index=column_index,
                        metric_name=metric_name,
                        unit=_unit_from_header(header),
                    )
                )
    return mapping


def _header_role(header: str) -> str | None:
    if header in {
        "from",
        "from date",
        "date from",
        "start date",
        "period start",
        "period start date",
        "billing start",
        "billing start date",
        "billing period start",
        "billing period start date",
    }:
        return "period_start"
    if header in {
        "to",
        "to date",
        "date to",
        "end date",
        "period end",
        "period end date",
        "billing end",
        "billing end date",
        "billing period end",
        "billing period end date",
    }:
        return "period_end"
    if header in {
        "period",
        "billing period",
        "bill period",
        "month",
        "billing month",
        "bill month",
        "date",
    }:
        return "period_start"
    if header in {"financial year", "fiscal year", "fy", "financial year fy"}:
        return "financial_year"
    if header in {"unit", "units", "uom", "unit of measure", "measurement unit"}:
        return "unit"
    if header in {
        "metric",
        "metric name",
        "parameter",
        "particular",
        "particulars",
        "indicator",
    }:
        return "metric_name"
    if header in {
        "value",
        "metric value",
        "numeric value",
        "quantity",
        "consumption value",
    }:
        return "metric_value"
    return None


def _known_metric_name(header: str) -> str | None:
    if "grid" in header and any(
        word in header for word in ("supply", "import", "consumption", "energy", "kwh")
    ):
        return "grid_supply_kwh"
    if "solar" in header and any(
        word in header for word in ("power", "generation", "energy", "kwh")
    ):
        return "solar_power_kwh"
    if "total" in header and any(
        word in header for word in ("consumption", "energy", "usage", "kwh")
    ):
        return "total_consumption"
    if any(
        phrase in header
        for phrase in (
            "units consumed",
            "energy consumption",
            "electricity consumption",
            "consumption kwh",
            "kwh consumption",
        )
    ):
        return "consumption_kwh"
    return None


def _is_known_header(header: str) -> bool:
    normalized = _normalize_header(header)
    return bool(
        normalized and (_header_role(normalized) or _known_metric_name(normalized))
    )


def _mapping_is_usable(mapping: SpreadsheetHeaderMapping) -> bool:
    return bool(
        mapping.metric_columns
        or (
            mapping.metric_name_column is not None
            and mapping.metric_value_column is not None
        )
    )


def _validate_mapping(mapping: SpreadsheetHeaderMapping, column_count: int) -> None:
    indexes = [
        mapping.metric_name_column,
        mapping.metric_value_column,
        mapping.period_start_column,
        mapping.period_end_column,
        mapping.financial_year_column,
        mapping.unit_column,
        *(metric.column_index for metric in mapping.metric_columns),
    ]
    if any(index is not None and index >= column_count for index in indexes):
        raise ValueError(
            "Column mapping references a column outside the spreadsheet header."
        )
    if any(index is not None and index < 0 for index in indexes):
        raise ValueError("Column mapping contains an invalid negative column index.")
    if not _mapping_is_usable(mapping):
        raise ValueError("No metric columns were mapped in the spreadsheet.")
    if len({metric.column_index for metric in mapping.metric_columns}) != len(
        mapping.metric_columns
    ):
        raise ValueError("Column mapping contains duplicate metric columns.")
    if (mapping.metric_name_column is None) != (mapping.metric_value_column is None):
        raise ValueError(
            "Both metric name and metric value columns are required for row-based data."
        )


async def _map_headers_with_groq(headers: list[str]) -> SpreadsheetHeaderMapping:
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        raise ValueError(
            "This spreadsheet layout is not recognized. Set GROQ_API_KEY to "
            "enable header mapping."
        )

    from langchain_groq import ChatGroq

    mapper = ChatGroq(
        model=os.getenv("GROQ_MODEL", "openai/gpt-oss-20b"),
        api_key=api_key,
        temperature=0,
    ).with_structured_output(SpreadsheetHeaderMapping)
    result = await mapper.ainvoke(
        [
            SystemMessage(content=SPREADSHEET_MAPPING_SYSTEM_PROMPT),
            HumanMessage(content=spreadsheet_mapping_input(headers)),
        ]
    )
    return SpreadsheetHeaderMapping.model_validate(result)


def _get_column_value(row: tuple, column_index: int | None):
    if column_index is None or column_index >= len(row):
        return None
    return row[column_index]


def _row_reference(sheet_name: str, row: tuple, row_number: int) -> str:
    populated_columns = [
        index + 1
        for index, value in enumerate(row)
        if value is not None and str(value).strip()
    ]
    if not populated_columns:
        raise ValueError(
            "Cannot create a source reference for an empty spreadsheet row."
        )
    first_column = get_column_letter(min(populated_columns))
    last_column = get_column_letter(max(populated_columns))
    if re.fullmatch(r"[A-Za-z0-9_]+", sheet_name):
        escaped_sheet = sheet_name
    else:
        escaped_sheet = "'" + sheet_name.replace("'", "''") + "'"
    start = f"{first_column}{row_number}"
    end = f"{last_column}{row_number}"
    cell_range = start if start == end else f"{start}:{end}"
    return f"{escaped_sheet}!{cell_range}"


def _cell_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _metric_name(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")
    if not normalized:
        raise ValueError("A metric name could not be determined from the source.")
    return normalized[:100]


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.casefold()).strip()


def _unit_from_header(header: str) -> str | None:
    match = _UNIT_PATTERN.search(header)
    if not match:
        return None
    unit = match.group(0).casefold()
    return {
        "kwh": "kWh",
        "mwh": "MWh",
        "wh": "Wh",
        "kw": "kW",
        "mw": "MW",
        "inr": "INR",
        "rupee": "INR",
        "rupees": "INR",
        "liter": "liters",
        "liters": "liters",
        "litre": "liters",
        "litres": "liters",
        "kg": "kg",
        "tonne": "tonnes",
        "tonnes": "tonnes",
        "ton": "tonnes",
        "tons": "tonnes",
    }.get(unit, unit)


def _to_decimal(value) -> Decimal | None:
    if value is None or isinstance(value, (bool, date, datetime)):
        return None
    try:
        if isinstance(value, Decimal):
            parsed = value
        elif isinstance(value, (int, float)):
            parsed = Decimal(str(value))
        else:
            text = str(value).strip()
            if not text:
                return None
            negative = text.startswith("(") and text.endswith(")")
            text = text.strip("()")
            text = re.sub(
                r"(?i)(₹|\$|INR|Rs\.?|USD|kWh|MWh|Wh|kg|liters?|litres?)",
                "",
                text,
            )
            text = text.replace(",", "").replace(" ", "")
            parsed = Decimal(text)
            if negative:
                parsed = -parsed
    except (InvalidOperation, ValueError):
        return None

    if not parsed.is_finite():
        return None
    _validate_decimal_precision(parsed)
    return parsed


def _validate_decimal_precision(value: Decimal) -> None:
    normalized = value.normalize()
    parts = normalized.as_tuple()
    scale = max(0, -parts.exponent)
    integer_digits = max(0, len(parts.digits) + parts.exponent)
    if scale > 12 or integer_digits > 26:
        raise ValueError("Metric value exceeds the supported exact numeric precision.")


def _financial_year(value: str | None, period_start: str | None) -> str | None:
    if value:
        match = re.search(r"(?<!\d)(\d{4})\s*[-/]\s*(\d{2,4})(?!\d)", value)
        if match:
            start_year = int(match.group(1))
            end_year = match.group(2)[-2:]
            return f"FY {start_year}-{end_year}"
        if value.casefold().startswith("fy "):
            return value
        return value

    parsed_date = _parse_period_date(period_start)
    if parsed_date is None:
        return None
    start_year = parsed_date.year if parsed_date.month >= 4 else parsed_date.year - 1
    return f"FY {start_year}-{(start_year + 1) % 100:02d}"


def _parse_period_date(value: str | None) -> date | None:
    if not value:
        return None
    for date_format in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%Y/%m/%d",
        "%B %Y",
        "%b %Y",
        "%Y-%m",
    ):
        try:
            return datetime.strptime(value, date_format).date()
        except ValueError:
            continue
    return None
