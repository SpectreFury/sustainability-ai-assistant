"""Convert uploaded PDFs and workbooks into source-located segments."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import PurePath
from typing import Iterable

import pymupdf
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from schemas.extraction import Locator, Segment
from services.text_sanitization import sanitize_text

ROWS_PER_SEGMENT = 60
MIN_PAGE_CHARS = 40
IMAGE_DPI = 144


class UnsupportedDocumentError(ValueError):
    """A document is unsupported, corrupt, or cannot be parsed."""


@dataclass
class ParsedDocument:
    filename: str
    kind: str
    page_count: int | None
    segments: list[Segment]
    page_images: dict[int, bytes] = field(default_factory=dict)


def kind_for(filename: str) -> str:
    suffix = PurePath(filename).suffix.casefold()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".xlsx":
        return "xlsx"
    raise UnsupportedDocumentError("Only PDF and XLSX documents are supported.")


def parse_document(
    data: bytes,
    filename: str,
    *,
    all_page_images: bool = False,
) -> ParsedDocument:
    kind = kind_for(filename)
    if not data:
        raise UnsupportedDocumentError("The uploaded document is empty.")
    try:
        if kind == "pdf":
            return _parse_pdf(data, filename, all_page_images)
        return _parse_xlsx(data, filename)
    except UnsupportedDocumentError:
        raise
    except Exception as error:
        raise UnsupportedDocumentError(
            f"Could not read {filename}: {type(error).__name__}."
        ) from error


def _parse_pdf(data: bytes, filename: str, all_page_images: bool) -> ParsedDocument:
    document = pymupdf.open(stream=data, filetype="pdf")
    try:
        if document.needs_pass:
            raise UnsupportedDocumentError(f"{filename} is password protected.")
        segments: list[Segment] = []
        images: dict[int, bytes] = {}
        for index, page in enumerate(document):
            page_number = index + 1
            text = sanitize_text(page.get_text("text", sort=True).strip()) or ""
            scanned = len("".join(text.split())) < MIN_PAGE_CHARS
            if scanned or all_page_images:
                images[page_number] = page.get_pixmap(dpi=IMAGE_DPI).tobytes("png")
            segments.append(
                Segment(
                    ordinal=index,
                    locator=Locator(page=page_number),
                    text=text,
                    scanned=scanned,
                )
            )
        if not segments:
            raise UnsupportedDocumentError(f"{filename} contains no pages.")
        return ParsedDocument(filename, "pdf", document.page_count, segments, images)
    finally:
        document.close()


def _cell_value(value: object) -> float | str:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        if isinstance(value, float) and not (float("-inf") < value < float("inf")):
            return str(value)
        return float(value)
    return sanitize_text(str(value).strip()) or ""


def _render(value: float | str) -> str:
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return value


def _sheet_segments(
    sheet_name: str,
    rows: Iterable[tuple[int, list[tuple[str, float | str, str | None]]]],
    start_ordinal: int,
) -> list[Segment]:
    segments: list[Segment] = []
    lines: list[str] = []
    cells: dict[str, float | str] = {}
    first_row: int | None = None
    last_row: int | None = None
    max_column = 1

    def flush() -> None:
        nonlocal lines, cells, first_row, last_row, max_column
        if not lines or first_row is None or last_row is None:
            return
        cell_range = f"A{first_row}:{get_column_letter(max_column)}{last_row}"
        header = f"Sheet '{sheet_name}', rows {first_row}-{last_row}"
        segments.append(
            Segment(
                ordinal=start_ordinal + len(segments),
                locator=Locator(sheet=sheet_name, cell_range=cell_range),
                text=header + "\n" + "\n".join(lines),
                cells=cells,
            )
        )
        lines, cells, first_row, last_row, max_column = [], {}, None, None, 1

    for row_number, row in rows:
        if not row:
            continue
        if first_row is None:
            first_row = row_number
        last_row = row_number
        parts: list[str] = []
        for reference, value, formula in row:
            cells[reference] = value
            column_name = re.match(r"[A-Z]+", reference)
            if column_name:
                max_column = max(max_column, _column_number(column_name.group(0)))
            rendered = f"{reference}={_render(value)}"
            if formula:
                rendered += f" (formula {formula})"
            parts.append(rendered)
        lines.append(" | ".join(parts))
        if len(lines) >= ROWS_PER_SEGMENT:
            flush()
    flush()
    return segments


def _column_number(column: str) -> int:
    result = 0
    for char in column:
        result = result * 26 + ord(char) - ord("A") + 1
    return result


def _parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    values_workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    formulas_workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    try:
        segments: list[Segment] = []
        for sheet_name in values_workbook.sheetnames:
            values_sheet = values_workbook[sheet_name]
            formula_rows = formulas_workbook[sheet_name].iter_rows()

            def rows(
                sheet=values_sheet,
                formula_row_iter=formula_rows,
            ) -> Iterable[tuple[int, list[tuple[str, float | str, str | None]]]]:
                for row_number, value_row in enumerate(sheet.iter_rows(), start=1):
                    formula_row = next(formula_row_iter, ())
                    row_cells: list[tuple[str, float | str, str | None]] = []
                    for column_number, cell in enumerate(value_row, start=1):
                        raw_formula_value = (
                            formula_row[column_number - 1].value
                            if column_number <= len(formula_row)
                            else None
                        )
                        formula = (
                            sanitize_text(raw_formula_value) or ""
                            if isinstance(raw_formula_value, str)
                            and raw_formula_value.startswith("=")
                            else None
                        )
                        value = cell.value
                        if value is None and formula is None:
                            continue
                        reference = f"{get_column_letter(column_number)}{row_number}"
                        row_cells.append(
                            (reference, _cell_value(value) if value is not None else "", formula)
                        )
                    yield row_number, row_cells

            segments.extend(_sheet_segments(sheet_name, rows(), len(segments)))
        if not segments:
            raise UnsupportedDocumentError(f"{filename} has no cell values.")
        return ParsedDocument(filename, "xlsx", None, segments)
    finally:
        values_workbook.close()
        formulas_workbook.close()
