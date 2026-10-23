import os
import re
import hashlib
from calendar import monthrange
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePath
from typing import Any, Literal
from uuid import UUID, uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from pypdf import PdfReader

from prompts.pdf_extraction import PDF_SYSTEM_PROMPT, PDF_USER_PROMPT
from services.fiscal_year import (
    financial_year_for_period,
    financial_year_label,
)
from services.llm_observability import (
    capture_content_enabled,
    langfuse_observation,
    update_observation,
)
from services.text_sanitization import sanitize_json, sanitize_text


class EvidenceCandidate(BaseModel):
    evidence_kind: Literal["pdf_page", "spreadsheet_cell", "calculation"]
    page_number: int | None = None
    printed_page_label: str | None = None
    sheet_name: str | None = None
    cell_reference: str | None = None
    range_reference: str | None = None
    quote: str | None = None
    bounding_box: dict[str, float] | None = None
    source_metadata: dict[str, Any] = Field(default_factory=dict)


class FactCandidate(BaseModel):
    id: UUID = Field(default_factory=uuid4)
    metric_name: str
    fact_kind: str = "metric"
    aggregation_role: str = "observation"
    raw_value: str | None = None
    metric_value: Decimal | None = None
    value_text: str | None = None
    value_date: date | None = None
    unit: str | None = None
    currency_code: str | None = None
    period_start: date | None = None
    period_end: date | None = None
    period_grain: str | None = None
    financial_year: str | None = None
    financial_year_start_year: int | None = None
    description: str | None = None
    dimensions: dict[str, Any] = Field(default_factory=dict)
    confidence: Decimal | None = None
    value_origin: Literal[
        "reported", "source_calculated", "system_calculated", "estimated"
    ] = "reported"
    calculation_method: str | None = None
    derived_from: list[UUID] = Field(default_factory=list)
    source_reference: str
    evidence: EvidenceCandidate


class ExtractedDocument(BaseModel):
    document_kind: Literal[
        "utility_bill", "spreadsheet", "operational_record", "policy", "unknown"
    ]
    entity_name: str | None = None
    facts: list[FactCandidate] = Field(default_factory=list)
    source_evidence: list[EvidenceCandidate] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)


class PdfFactResponse(BaseModel):
    metric_name: str = Field(
        description="Stable snake_case fact name, not a prose sentence"
    )
    fact_kind: Literal[
        "metric",
        "financial_amount",
        "meter_reading",
        "record",
        "policy_commitment",
        "metadata",
    ] = "metric"
    aggregation_role: Literal[
        "observation", "component", "subtotal", "reported_total", "commitment"
    ] = "observation"
    value_type: Literal["numeric", "text", "date"] = "numeric"
    raw_value: str = Field(description="Value exactly as printed in the document")
    unit: str | None = None
    currency_code: str | None = None
    event_date: str | None = None
    period_start: str | None = None
    period_end: str | None = None
    period_grain: str | None = None
    description: str | None = None
    dimensions: dict[str, Any] = Field(default_factory=dict)
    page_number: int = Field(ge=1)
    source_quote: str = Field(
        min_length=1,
        description="Short exact source text sufficient to locate and verify the fact",
    )
    confidence: float = Field(default=0.8, ge=0, le=1)


class PdfExtractionResponse(BaseModel):
    document_kind: Literal[
        "utility_bill", "operational_record", "policy", "unknown"
    ]
    entity_name: str | None = None
    facts: list[PdfFactResponse] = Field(default_factory=list)


_MONTHS = {
    name.casefold(): index
    for index, name in enumerate(
        (
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        ),
        start=1,
    )
}
_MONTHS.update({name[:3]: number for name, number in list(_MONTHS.items())})
_YEAR_PATTERN = re.compile(r"\b(19\d{2}|20\d{2})\b")
_YEAR_CONTEXT_PATTERN = re.compile(
    r"(?:full|calendar|financial|fiscal)?\s*year\s*(?:of\s*)?(19\d{2}|20\d{2})",
    re.IGNORECASE,
)
_MONTH_RANGE_PATTERN = re.compile(
    r"^\s*([a-z]{3,9})\s*[-–—]\s*([a-z]{3,9})\s*(19\d{2}|20\d{2})?\s*$",
    re.IGNORECASE,
)
_FINANCIAL_YEAR_PATTERN = re.compile(
    r"(?:FY\s*)?(19\d{2}|20\d{2})\s*[-/]\s*(\d{2}|\d{4})",
    re.IGNORECASE,
)
_ADDITIVE_METRICS = {
    "grid_electricity_consumption",
    "solar_generation",
    "wind_generation",
    "renewable_generation",
    "renewable_energy_savings",
    "total_energy_consumption",
    "energy_generation",
    "grid_cost",
    "renewable_savings",
    "net_cost",
    "co2_avoided",
    "quantity",
}


async def extract_document_facts(
    filename: str,
    content: bytes,
    trace_context: dict[str, str] | None = None,
) -> ExtractedDocument:
    extension = PurePath(filename).suffix.casefold()
    if extension == ".xlsx":
        return extract_spreadsheet_facts(content)
    if extension == ".pdf":
        return await extract_pdf_facts(content, trace_context=trace_context)
    raise ValueError("Unsupported file type. Only .xlsx and .pdf are supported.")


def extract_spreadsheet_facts(content: bytes) -> ExtractedDocument:
    formulas = load_workbook(BytesIO(content), read_only=False, data_only=False)
    values = load_workbook(BytesIO(content), read_only=False, data_only=True)
    facts: list[FactCandidate] = []
    warnings: list[str] = []
    year_context = _workbook_year_context(formulas)

    try:
        for formula_sheet in formulas.worksheets:
            value_sheet = values[formula_sheet.title]
            header_row, headers = _find_table_header(formula_sheet)
            if header_row is None:
                continue

            header_map = {
                index: _metric_definition(header)
                for index, header in enumerate(headers)
            }
            metric_columns = {
                index: definition
                for index, definition in header_map.items()
                if definition is not None
            }
            if not metric_columns:
                continue

            sheet_year = _sheet_year(formula_sheet, year_context)
            header_end = get_column_letter(len(headers))
            for row_number in range(header_row + 1, formula_sheet.max_row + 1):
                formula_row = next(
                    formula_sheet.iter_rows(
                        min_row=row_number,
                        max_row=row_number,
                        max_col=len(headers),
                    )
                )
                value_row = next(
                    value_sheet.iter_rows(
                        min_row=row_number,
                        max_row=row_number,
                        max_col=len(headers),
                    )
                )
                if not any(cell.value is not None for cell in formula_row):
                    continue

                row_values = [cell.value for cell in value_row]
                formula_values = [cell.value for cell in formula_row]
                row_label = _first_text(row_values)
                period_start, period_end, period_grain, period_source = (
                    _row_period(
                        headers,
                        row_values,
                        row_label,
                        sheet_year,
                        formula_sheet.title,
                    )
                )
                aggregation_role = _row_aggregation_role(
                    row_label, period_grain, formula_sheet.title
                )
                unit_column = next(
                    (
                        index
                        for index, header in enumerate(headers)
                        if _period_column_role(header) == "unit"
                    ),
                    None,
                )
                row_unit = (
                    _display_unit(row_values[unit_column])
                    if unit_column is not None and unit_column < len(row_values)
                    else None
                )
                row_metric = _row_label_metric(row_label, row_unit)

                for column_index, definition in metric_columns.items():
                    value_cell = value_row[column_index]
                    formula_cell = formula_row[column_index]
                    cached_value = value_cell.value
                    formula_value = formula_cell.value
                    is_formula = (
                        isinstance(formula_value, str)
                        and formula_value.startswith("=")
                    )
                    formula_without_cached_result = is_formula and cached_value is None
                    source_value = (
                        formula_value if formula_without_cached_result else cached_value
                    )
                    if source_value is None or _is_empty_marker(source_value):
                        continue

                    metric_name, header_unit, currency_code = definition
                    adjacent_unit = _adjacent_unit(row_values, column_index)
                    cell_unit_context = adjacent_unit or (row_unit if not header_unit else None)
                    if row_metric is not None:
                        row_override = _row_metric_override(
                            row_label or "",
                            cell_unit_context,
                            header_unit,
                        )
                        if row_override is not None:
                            metric_name, _, row_currency = row_override
                            currency_code = row_currency or currency_code
                    if (
                        row_label
                        and cell_unit_context
                        and header_unit
                        and _normalize_unit_value(cell_unit_context)
                        != _normalize_unit_value(header_unit)
                    ):
                        row_override = _row_metric_override(
                            row_label,
                            cell_unit_context,
                            header_unit,
                        )
                        if row_override is None:
                            metric_name = _slug(row_label)
                            currency_code = None
                    if formula_without_cached_result:
                        metric_value, parsed_unit = None, None
                    else:
                        metric_value, parsed_unit = _spreadsheet_decimal(source_value)
                    unit = parsed_unit or adjacent_unit or header_unit or row_unit
                    if currency_code is None and _normalize_unit_value(unit) == "inr":
                        currency_code = "INR"
                    if metric_value is None:
                        # Preserve relevant text facts, but don't turn labels into numbers.
                        if formula_without_cached_result:
                            value_date = None
                            value_text = None
                        elif isinstance(source_value, (date, datetime)):
                            value_date = (
                                source_value.date()
                                if isinstance(source_value, datetime)
                                else source_value
                            )
                            value_text = None
                        else:
                            continue
                    else:
                        value_date = None
                        value_text = None

                    cell_reference = value_cell.coordinate
                    header_cell = formula_sheet.cell(header_row, column_index + 1)
                    label_cell = formula_sheet.cell(row_number, 1)
                    label_text = _cell_to_text(label_cell.value)
                    quote = (
                        f"{label_text}; {header_cell.value}: {_cell_to_text(source_value)}"
                        if label_text
                        else f"{header_cell.value}: {_cell_to_text(source_value)}"
                    )
                    dimensions = _row_dimensions(
                        headers,
                        row_values,
                        period_grain,
                        row_label,
                    )
                    if period_source:
                        dimensions["period_source"] = period_source
                    if is_formula:
                        dimensions["source_formula"] = str(formula_value)

                    source_reference = (
                        f"{_excel_sheet_reference(formula_sheet.title)}!{cell_reference}"
                    )
                    financial_year_start = financial_year_for_period(
                        period_start, period_end
                    )
                    fact = FactCandidate(
                        metric_name=metric_name,
                        fact_kind=(
                            "formula"
                            if formula_without_cached_result
                            else "financial_amount"
                            if currency_code
                            else "metric"
                        ),
                        aggregation_role=aggregation_role,
                        raw_value=_cell_to_text(source_value),
                        metric_value=metric_value,
                        value_text=value_text,
                        value_date=value_date,
                        unit=unit,
                        currency_code=currency_code,
                        period_start=period_start,
                        period_end=period_end,
                        period_grain=period_grain,
                        financial_year=(
                            financial_year_label(financial_year_start)
                            if financial_year_start is not None
                            else None
                        ),
                        financial_year_start_year=financial_year_start,
                        description=(
                            label_text
                            or "Formula result is not cached in the workbook."
                            if formula_without_cached_result
                            else label_text
                        ),
                        dimensions=dimensions,
                        confidence=Decimal("1.000"),
                        value_origin=("source_calculated" if is_formula else "reported"),
                        source_reference=source_reference,
                        evidence=EvidenceCandidate(
                            evidence_kind="spreadsheet_cell",
                            sheet_name=formula_sheet.title,
                            cell_reference=cell_reference,
                            range_reference=f"A{header_row}:{header_end}{row_number}",
                            quote=quote,
                            source_metadata={
                                "header_cell": header_cell.coordinate,
                                "row_label_cell": label_cell.coordinate,
                                "header": _cell_to_text(header_cell.value),
                                "row_label": label_text,
                                "year_context": sheet_year,
                                "source_formula": str(formula_value) if is_formula else None,
                            },
                        ),
                    )
                    facts.append(fact)

                # Some source sheets contain compact label/value/unit rows inside
                # otherwise tabular layouts (for example, a CO2 value followed by
                # "Tonnes"). Preserve those as separate, precisely cited facts.
                if row_label and row_metric is not None:
                    for column_index in range(len(row_values) - 1):
                        if column_index in metric_columns:
                            continue
                        source_value = value_row[column_index].value
                        if (
                            isinstance(source_value, str)
                            and source_value.startswith("=")
                        ):
                            continue
                        adjacent_unit_raw = value_row[column_index + 1].value
                        adjacent_unit = _adjacent_unit(row_values, column_index)
                        if adjacent_unit is None:
                            continue
                        metric_value, parsed_unit = _spreadsheet_decimal(source_value)
                        if metric_value is None:
                            continue
                        override = _row_metric_override(
                            row_label,
                            adjacent_unit,
                            adjacent_unit,
                        )
                        if override is None:
                            continue
                        metric_name, unit, currency_code = override
                        value_cell = value_row[column_index]
                        label_cell = formula_sheet.cell(row_number, 1)
                        source_reference = (
                            f"{_excel_sheet_reference(formula_sheet.title)}!"
                            f"{value_cell.coordinate}"
                        )
                        facts.append(
                            FactCandidate(
                                metric_name=metric_name,
                                fact_kind=(
                                    "financial_amount" if currency_code else "metric"
                                ),
                                aggregation_role=aggregation_role,
                                raw_value=_cell_to_text(source_value),
                                metric_value=metric_value,
                                unit=parsed_unit or unit,
                                currency_code=currency_code,
                                period_start=period_start,
                                period_end=period_end,
                                period_grain=period_grain,
                                financial_year=(
                                    financial_year_label(
                                        financial_year_for_period(
                                            period_start, period_end
                                        )
                                    )
                                    if financial_year_for_period(
                                        period_start, period_end
                                    )
                                    is not None
                                    else None
                                ),
                                financial_year_start_year=financial_year_for_period(
                                    period_start, period_end
                                ),
                                description=row_label,
                                dimensions={
                                    "row_label": row_label,
                                    "paired_unit_cell": value_row[
                                        column_index + 1
                                    ].coordinate,
                                },
                                confidence=Decimal("1.000"),
                                value_origin="reported",
                                source_reference=source_reference,
                                evidence=EvidenceCandidate(
                                    evidence_kind="spreadsheet_cell",
                                    sheet_name=formula_sheet.title,
                                    cell_reference=value_cell.coordinate,
                                    range_reference=(
                                        f"{value_cell.coordinate}:"
                                        f"{value_row[column_index + 1].coordinate}"
                                    ),
                                    quote=(
                                        f"{row_label}: {_cell_to_text(source_value)} "
                                        f"{adjacent_unit}"
                                    ),
                                    source_metadata={
                                        "row_label_cell": label_cell.coordinate,
                                        "paired_unit_cell": value_row[
                                            column_index + 1
                                        ].coordinate,
                                    },
                                ),
                            )
                        )

            if sheet_year is None and any(
                fact.period_grain == "month"
                and fact.evidence.sheet_name == formula_sheet.title
                for fact in facts
            ):
                warnings.append(
                    f"{formula_sheet.title}: month rows have no reliable year in the workbook."
                )
    finally:
        formulas.close()
        values.close()

    if not facts:
        raise ValueError("No supported numeric tables were found in the spreadsheet.")

    warnings.extend(_add_period_reconciliations(facts))
    warnings.extend(_financial_year_coverage_warnings(facts))
    return ExtractedDocument(
        document_kind="spreadsheet",
        facts=facts,
        warnings=warnings,
    )


async def extract_pdf_facts(
    content: bytes,
    trace_context: dict[str, str] | None = None,
) -> ExtractedDocument:
    reader = PdfReader(BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are not supported.")
    if not reader.pages:
        raise ValueError("The PDF contains no pages.")

    page_texts = [
        sanitize_text(page.extract_text() or "") or "" for page in reader.pages
    ]
    source_evidence = [
        EvidenceCandidate(
            evidence_kind="pdf_page",
            page_number=page_number,
            quote=page_text[:20000],
            source_metadata={
                "content_kind": "page_text",
                "text_length": len(page_text),
                "truncated": len(page_text) > 20000,
            },
        )
        for page_number, page_text in enumerate(page_texts, start=1)
        if page_text.strip()
    ]
    extraction = await _invoke_pdf_model(content, page_texts, trace_context)
    facts: list[FactCandidate] = []
    warnings: list[str] = []

    for fact in extraction.facts:
        if fact.page_number > len(page_texts):
            warnings.append(
                f"Ignored a fact with invalid PDF page {fact.page_number}."
            )
            continue
        raw_value = sanitize_text(fact.raw_value) or ""
        source_quote = sanitize_text(fact.source_quote) or ""
        metric_name = sanitize_text(fact.metric_name) or "unknown"
        numeric_value: Decimal | None = None
        text_value: str | None = None
        date_value: date | None = None
        if fact.value_type == "numeric":
            numeric_value = _parse_decimal(raw_value)
            if numeric_value is None:
                warnings.append(
                    f"Ignored non-numeric value for {metric_name} on page "
                    f"{fact.page_number}."
                )
                continue
        elif fact.value_type == "date":
            date_value = _parse_date(
                sanitize_text(fact.event_date) or raw_value
            )
            if date_value is None:
                text_value = raw_value
                warnings.append(
                    f"Could not normalize date for {metric_name} on page "
                    f"{fact.page_number}; raw value was retained as text."
                )
            else:
                numeric_value = None
        else:
            text_value = raw_value

        period_start = _parse_date(sanitize_text(fact.period_start))
        period_end = _parse_date(sanitize_text(fact.period_end))
        if fact.event_date and date_value is None:
            date_value = _parse_date(sanitize_text(fact.event_date))
        fy_start_year = financial_year_for_period(period_start, period_end)
        page_text = page_texts[fact.page_number - 1]
        quote_verified = _quote_is_present(source_quote, page_text)
        if quote_verified is False:
            warnings.append(
                f"Source quote for {metric_name} on page {fact.page_number} "
                "did not match extracted PDF text; review is recommended."
            )

        source_reference = f"PDF page {fact.page_number}"
        facts.append(
            FactCandidate(
                metric_name=_slug(metric_name),
                fact_kind=fact.fact_kind,
                aggregation_role=fact.aggregation_role,
                raw_value=raw_value,
                metric_value=numeric_value,
                value_text=sanitize_text(text_value),
                value_date=date_value,
                unit=sanitize_text(fact.unit),
                currency_code=sanitize_text(fact.currency_code),
                period_start=period_start,
                period_end=period_end,
                period_grain=fact.period_grain or _period_grain(period_start, period_end),
                financial_year=(
                    financial_year_label(fy_start_year)
                    if fy_start_year is not None
                    else None
                ),
                financial_year_start_year=fy_start_year,
                description=sanitize_text(fact.description),
                dimensions=sanitize_json(fact.dimensions),
                confidence=Decimal(str(fact.confidence)),
                value_origin="reported",
                source_reference=source_reference,
                evidence=EvidenceCandidate(
                    evidence_kind="pdf_page",
                    page_number=fact.page_number,
                    quote=source_quote,
                    source_metadata={"quote_verified": quote_verified},
                ),
            )
        )

    if not facts:
        raise ValueError("No verifiable facts were extracted from the PDF.")
    return ExtractedDocument(
        document_kind=extraction.document_kind,
        entity_name=sanitize_text(extraction.entity_name),
        facts=facts,
        source_evidence=source_evidence,
        warnings=warnings,
    )


async def _invoke_pdf_model(
    content: bytes,
    page_texts: list[str],
    trace_context: dict[str, str] | None,
) -> PdfExtractionResponse:
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not api_key:
        raise ValueError("Set GEMINI_API_KEY to enable PDF extraction.")

    from langchain_core.messages import HumanMessage, SystemMessage
    from langchain_google_genai import ChatGoogleGenerativeAI

    model_name = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
    model = ChatGoogleGenerativeAI(
        model=model_name,
        api_key=api_key,
    ).with_structured_output(PdfExtractionResponse, include_raw=True)
    text_context = "\n\n".join(
        f"--- PDF page {index + 1} ---\n{text[:16000]}"
        for index, text in enumerate(page_texts)
        if text.strip()
    )
    user_text = PDF_USER_PROMPT
    if text_context:
        user_text += f"\n\nExtracted page text for exact citations:\n{text_context}"

    trace_metadata = {
        **(trace_context or {}),
        "document_format": "pdf",
        "page_count": str(len(page_texts)),
        "content_sha256": hashlib.sha256(content).hexdigest(),
        "prompt_version": "pdf-facts-v1",
    }
    trace_input: dict[str, Any] = {
        "document_id": trace_metadata.get("document_id"),
        "extraction_run_id": trace_metadata.get("extraction_run_id"),
        "page_count": len(page_texts),
        "content_sha256": trace_metadata["content_sha256"],
        "page_text_char_count": sum(len(text) for text in page_texts),
        "prompt_version": trace_metadata["prompt_version"],
    }
    generation_input: dict[str, Any] = {
        "document_format": "pdf",
        "page_count": len(page_texts),
        "prompt_version": trace_metadata["prompt_version"],
        "system_prompt": PDF_SYSTEM_PROMPT,
        "user_prompt": PDF_USER_PROMPT,
    }
    if capture_content_enabled():
        # Trace extracted text and structured results, not the binary PDF payload.
        trace_input["extracted_page_text"] = text_context[:60000]
        generation_input["extracted_page_text"] = text_context[:60000]

    with langfuse_observation(
        as_type="span",
        name="sustainability-document-pdf-extraction",
        input=trace_input,
        metadata=trace_metadata,
    ) as trace:
        with langfuse_observation(
            as_type="generation",
            name="gemini-structured-pdf-fact-extraction",
            model=model_name,
            input=generation_input,
            metadata={
                "provider": "google",
                "prompt_version": trace_metadata["prompt_version"],
            },
        ) as generation:
            try:
                raw_result = await model.ainvoke(
                    [
                        SystemMessage(content=PDF_SYSTEM_PROMPT),
                        HumanMessage(
                            content=[
                                {"type": "text", "text": user_text},
                                {
                                    "type": "media",
                                    "mime_type": "application/pdf",
                                    "data": content,
                                },
                            ]
                        ),
                    ]
                )
                if isinstance(raw_result, dict) and "parsed" in raw_result:
                    parsing_error = raw_result.get("parsing_error")
                    if parsing_error is not None:
                        raise ValueError("Gemini returned an invalid structured extraction.")
                    extraction = PdfExtractionResponse.model_validate(
                        raw_result.get("parsed")
                    )
                    raw_message = raw_result.get("raw")
                else:
                    extraction = PdfExtractionResponse.model_validate(raw_result)
                    raw_message = None

                usage_details = _langchain_token_usage(raw_message)
                generation_output = _pdf_trace_output(extraction)
                update_kwargs: dict[str, Any] = {"output": generation_output}
                if usage_details:
                    update_kwargs["usage_details"] = usage_details
                update_observation(generation, **update_kwargs)
                update_observation(
                    trace,
                    output={
                        "document_kind": extraction.document_kind,
                        "fact_count": len(extraction.facts),
                        "models": model_name,
                    },
                )
                return extraction
            except Exception as error:
                update_observation(
                    generation,
                    level="ERROR",
                    status_message=type(error).__name__,
                )
                update_observation(
                    trace,
                    level="ERROR",
                    status_message=type(error).__name__,
                )
                raise


def _langchain_token_usage(raw_message: Any) -> dict[str, int] | None:
    if raw_message is None:
        return None
    usage = getattr(raw_message, "usage_metadata", None)
    if not usage:
        response_metadata = getattr(raw_message, "response_metadata", None) or {}
        usage = response_metadata.get("token_usage") or {}
    if not isinstance(usage, dict):
        return None
    aliases = {
        "input": usage.get("input_tokens", usage.get("prompt_tokens")),
        "output": usage.get("output_tokens", usage.get("completion_tokens")),
        "total": usage.get("total_tokens"),
    }
    result = {key: int(value) for key, value in aliases.items() if value is not None}
    return result or None


def _pdf_trace_output(extraction: PdfExtractionResponse) -> dict[str, Any]:
    if capture_content_enabled():
        return extraction.model_dump(mode="json")
    return {
        "document_kind": extraction.document_kind,
        "fact_count": len(extraction.facts),
        "metric_names": sorted({fact.metric_name for fact in extraction.facts}),
        "cited_pages": sorted({fact.page_number for fact in extraction.facts}),
    }


def _find_table_header(worksheet) -> tuple[int | None, list[str]]:
    best_row: int | None = None
    best_headers: list[str] = []
    best_score = 0
    for row_number in range(1, min(worksheet.max_row or 1, 30) + 1):
        values = next(
            worksheet.iter_rows(
                min_row=row_number,
                max_row=row_number,
                max_col=worksheet.max_column or 1,
                values_only=True,
            ),
            (),
        )
        headers = [_cell_to_text(value) or "" for value in values]
        metric_count = sum(_metric_definition(header) is not None for header in headers)
        period_count = sum(_period_column_role(header) is not None for header in headers)
        score = metric_count * 10 + period_count * 3 + sum(bool(h) for h in headers)
        if metric_count and score > best_score:
            best_row = row_number
            best_headers = headers
            best_score = score
    return best_row, best_headers


def _metric_definition(header: str) -> tuple[str, str | None, str | None] | None:
    normalized = _normalize_header(header)
    if not normalized:
        return None

    aliases: tuple[tuple[tuple[str, ...], str], ...] = (
        (("average monthly grid consumption",), "average_monthly_grid_consumption"),
        (("average daily grid consumption",), "average_daily_grid_consumption"),
        (("total co2 emissions avoided", "co2 avoided"), "co2_avoided"),
        (("total renewable", "renewable generation"), "renewable_generation"),
        (("total consumption", "total energy consumption", "total kwh"), "total_energy_consumption"),
        (("grid supply", "grid electricity", "grid consumption", "grid kwh"), "grid_electricity_consumption"),
        (("solar annual generation", "solar generation", "solar power", "solar kwh"), "solar_generation"),
        (("wind annual generation", "wind generation", "wind power", "wind kwh"), "wind_generation"),
        (("annual generation", "energy generation"), "energy_generation"),
        (("renewable savings", "savings"), "renewable_savings"),
        (("grid cost", "actual grid cost"), "grid_cost"),
        (("net effective cost", "net cost"), "net_cost"),
        (("total cost", "cost amount"), "total_cost"),
        (("cost per kwh", "average effective cost"), "cost_per_kwh"),
        (("renewable %", "renewable percent", "renewable share"), "renewable_share_percent"),
        (("grid %", "grid share"), "grid_share_percent"),
        (("solar %", "solar share"), "solar_share_percent"),
        (("wind %", "wind share"), "wind_share_percent"),
        (("% of total", "percentage of total"), "share_percent"),
        (("capacity utilization factor", "utilization factor"), "capacity_utilization_percent"),
        (("solar capacity",), "solar_capacity"),
        (("wind capacity",), "wind_capacity"),
        (("solar average monthly",), "average_monthly_solar_generation"),
        (("wind average monthly",), "average_monthly_wind_generation"),
        (("cost reduction",), "cost_reduction"),
        (("quantity", "total quantity"), "quantity"),
        (("amount payable", "net payable amount", "amount due"), "amount_payable"),
        (("amount", "total amount"), "amount"),
        (("unit rate", "rate per"), "unit_rate"),
        (("value", "metric value"), "value"),
    )
    metric_name: str | None = None
    for patterns, candidate in aliases:
        if any(pattern in normalized for pattern in patterns):
            metric_name = candidate
            break
    if metric_name is None:
        return None

    unit, currency_code = _header_unit(header)
    return metric_name, unit, currency_code


def _period_column_role(header: str) -> str | None:
    normalized = _normalize_header(header)
    if normalized in {"month", "billing month", "date", "record date"}:
        return "month"
    if normalized in {"period", "billing period", "reporting period"}:
        return "period"
    if normalized in {"quarter", "financial quarter"}:
        return "quarter"
    if normalized in {"financial year", "fiscal year", "fy"}:
        return "financial_year"
    if normalized in {"unit", "units", "uom", "unit of measure"}:
        return "unit"
    return None


def _row_label_metric(
    row_label: str | None,
    row_unit: str | None,
) -> tuple[str, str | None, str | None] | None:
    if not row_label:
        return None
    normalized_label = _normalize_header(row_label)
    normalized_unit = _normalize_unit_value(row_unit)
    if normalized_label == "total" and normalized_unit in {"kwh", "mwh", "wh"}:
        return "total_energy_consumption", "kWh", None
    if "actual grid cost" in normalized_label and normalized_unit in {
        "kwh",
        "mwh",
        "wh",
    }:
        return "grid_electricity_consumption", "kWh", None
    if "renewable savings" in normalized_label and normalized_unit in {
        "kwh",
        "mwh",
        "wh",
    }:
        return "renewable_energy_savings", "kWh", None
    if "net effective cost" in normalized_label and normalized_unit in {
        "kwh",
        "mwh",
        "wh",
    }:
        return "total_energy_consumption", "kWh", None
    return _metric_definition(row_label)


def _row_metric_override(
    row_label: str,
    row_unit: str | None,
    header_unit: str | None,
) -> tuple[str, str | None, str | None] | None:
    candidate = _row_label_metric(row_label, row_unit)
    if candidate is None:
        return None
    metric_name, candidate_unit, currency_code = candidate
    observed_unit = _normalize_unit_value(row_unit) or _normalize_unit_value(header_unit)
    allowed_units = {
        "grid_electricity_consumption": {"kwh", "mwh", "wh"},
        "solar_generation": {"kwh", "mwh", "wh"},
        "wind_generation": {"kwh", "mwh", "wh"},
        "renewable_generation": {"kwh", "mwh", "wh"},
        "renewable_energy_savings": {"kwh", "mwh", "wh"},
        "total_energy_consumption": {"kwh", "mwh", "wh"},
        "energy_generation": {"kwh", "mwh", "wh"},
        "co2_avoided": {"kg", "tonnes"},
        "grid_cost": {"inr"},
        "renewable_savings": {"inr"},
        "net_cost": {"inr"},
        "total_cost": {"inr"},
        "cost_per_kwh": {"inr/kwh"},
        "solar_capacity": {"kwp", "kw"},
        "wind_capacity": {"kw"},
        "capacity_utilization_percent": {"%"},
        "renewable_share_percent": {"%"},
    }
    expected = allowed_units.get(metric_name)
    if expected and observed_unit not in expected:
        return None
    return metric_name, row_unit or candidate_unit or header_unit, currency_code


def _normalize_unit_value(value: Any) -> str | None:
    text = _cell_to_text(value)
    if not text:
        return None
    normalized = text.casefold().replace(" ", "")
    aliases = {
        "kwh": "kwh",
        "mwh": "mwh",
        "wh": "wh",
        "inr": "inr",
        "₹": "inr",
        "rs": "inr",
        "rupees": "inr",
        "%": "%",
        "percent": "%",
        "kg": "kg",
        "kgs": "kg",
        "tonne": "tonnes",
        "tonnes": "tonnes",
        "t": "tonnes",
        "mt": "tonnes",
        "kwp": "kwp",
        "kw": "kw",
        "inr/kwh": "inr/kwh",
        "₹/kwh": "inr/kwh",
        "rs/kwh": "inr/kwh",
    }
    return aliases.get(normalized, normalized)


def _display_unit(value: Any) -> str | None:
    normalized = _normalize_unit_value(value)
    return {
        "kwh": "kWh",
        "mwh": "MWh",
        "wh": "Wh",
        "inr": "INR",
        "%": "%",
        "kg": "kg",
        "tonnes": "tonnes",
        "kwp": "kWp",
        "kw": "kW",
        "inr/kwh": "INR/kWh",
    }.get(normalized, _cell_to_text(value))


def _adjacent_unit(row_values: list[Any], value_column: int) -> str | None:
    unit_column = value_column + 1
    if unit_column >= len(row_values):
        return None
    value = row_values[unit_column]
    normalized = _normalize_unit_value(value)
    known_units = {
        "kwh",
        "mwh",
        "wh",
        "inr",
        "%",
        "kg",
        "tonnes",
        "kwp",
        "kw",
        "inr/kwh",
        "trees",
        "cars/year",
        "m³",
        "liters",
    }
    if normalized not in known_units:
        return None
    return _display_unit(value)


def _row_period(
    headers: list[str],
    row_values: list[Any],
    row_label: str | None,
    year_context: int | None,
    sheet_name: str,
) -> tuple[date | None, date | None, str | None, str | None]:
    period_values: list[tuple[str, Any]] = []
    for index, header in enumerate(headers):
        role = _period_column_role(header)
        if role in {"month", "period", "quarter", "financial_year"}:
            value = row_values[index] if index < len(row_values) else None
            if value is not None:
                period_values.append((role, value))

    for role, value in period_values:
        parsed = _parse_period_value(value, year_context)
        if parsed[0] is not None:
            return (*parsed, f"column:{role}")

    if row_label:
        parsed = _parse_period_value(row_label, year_context)
        if parsed[0] is not None:
            return (*parsed, "row_label")

    if year_context is not None and "annual" in sheet_name.casefold():
        return (
            date(year_context, 1, 1),
            date(year_context, 12, 31),
            "year",
            "annual_sheet_context",
        )

    for value in row_values:
        text = _cell_to_text(value)
        if text:
            match = _YEAR_CONTEXT_PATTERN.search(text)
            if match:
                year = int(match.group(1))
                return date(year, 1, 1), date(year, 12, 31), "year", "row_context"

    if year_context is not None and row_label:
        normalized_label = _normalize_header(row_label)
        if normalized_label in {
            "annual",
            "annual total",
            "total",
            "full year",
            "calendar year",
        }:
            return (
                date(year_context, 1, 1),
                date(year_context, 12, 31),
                "year",
                "workbook_year_context",
            )
    return None, None, None, None


def _parse_period_value(
    value: Any,
    year_context: int | None,
) -> tuple[date | None, date | None, str | None]:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        last_day = monthrange(value.year, value.month)[1]
        return date(value.year, value.month, 1), date(value.year, value.month, last_day), "month"

    text = _cell_to_text(value)
    if not text:
        return None, None, None
    match_fy = _FINANCIAL_YEAR_PATTERN.search(text)
    if match_fy:
        start_year = int(match_fy.group(1))
        return date(start_year, 4, 1), date(start_year + 1, 3, 31), "financial_year"

    direct_date = _parse_date(text)
    if direct_date is not None:
        last_day = monthrange(direct_date.year, direct_date.month)[1]
        return (
            date(direct_date.year, direct_date.month, 1),
            date(direct_date.year, direct_date.month, last_day),
            "month",
        )

    year_match = _YEAR_PATTERN.search(text)
    year = int(year_match.group(1)) if year_match else year_context
    range_match = _MONTH_RANGE_PATTERN.match(text)
    if range_match and year is not None:
        start_month = _month_number(range_match.group(1))
        end_month = _month_number(range_match.group(2))
        if start_month and end_month:
            end_year = year + (1 if end_month < start_month else 0)
            grain = (
                "year"
                if start_month == 1 and end_month == 12 and end_year == year
                else "quarter"
                if end_month - start_month <= 2
                else "range"
            )
            return (
                date(year, start_month, 1),
                date(end_year, end_month, monthrange(end_year, end_month)[1]),
                grain,
            )

    if year is not None:
        month = _month_number(text)
        if month is not None:
            return date(year, month, 1), date(year, month, monthrange(year, month)[1]), "month"
        quarter = re.fullmatch(r"q([1-4])", text.strip(), re.IGNORECASE)
        if quarter:
            quarter_number = int(quarter.group(1))
            start_month = (quarter_number - 1) * 3 + 1
            end_month = start_month + 2
            return (
                date(year, start_month, 1),
                date(year, end_month, monthrange(year, end_month)[1]),
                "quarter",
            )
        if re.fullmatch(r"(?:annual|year|full year)\s*", text, re.IGNORECASE):
            return date(year, 1, 1), date(year, 12, 31), "year"

    if year_match and re.fullmatch(r"(?:19|20)\d{2}", text.strip()):
        year = int(text.strip())
        return date(year, 1, 1), date(year, 12, 31), "year"
    return None, None, None


def _row_aggregation_role(
    row_label: str | None,
    period_grain: str | None,
    sheet_name: str,
) -> str:
    normalized = _normalize_header(row_label or "")
    if period_grain in {"month", "day"}:
        return "observation"
    if period_grain in {"quarter", "range", "financial_year", "year"}:
        return "reported_total"
    if re.fullmatch(r"q[1-4]", normalized):
        return "reported_total"
    if normalized in {"annual", "annual total", "total", "full year", "calendar year"}:
        return "reported_total"
    if "annual" in sheet_name.casefold():
        return "reported_summary"
    return "reported_summary"


def _row_dimensions(
    headers: list[str],
    row_values: list[Any],
    period_grain: str | None,
    row_label: str | None,
) -> dict[str, Any]:
    dimensions: dict[str, Any] = {}
    if not headers or not row_values:
        return dimensions
    label_header = _normalize_header(headers[0])
    if row_label and label_header not in {
        "month",
        "date",
        "period",
        "quarter",
        "financial year",
        "fy",
    }:
        key = _slug(label_header) or "row_label"
        dimensions[key] = row_label
    elif row_label and period_grain == "quarter":
        dimensions["quarter"] = row_label
    return dimensions


def _add_period_reconciliations(facts: list[FactCandidate]) -> list[str]:
    reported_totals = [
        fact
        for fact in facts
        if fact.aggregation_role == "reported_total"
        and fact.period_grain in {"quarter", "year"}
        and fact.period_start is not None
        and fact.period_end is not None
        and fact.metric_name in _ADDITIVE_METRICS
    ]
    additions: list[FactCandidate] = []
    warnings: list[str] = []
    calculated_by_period: dict[tuple[str, str | None, date, date], FactCandidate] = {}

    for reported in reported_totals:
        assert reported.period_start is not None
        assert reported.period_end is not None
        expected_months = _months_in_period(reported.period_start, reported.period_end)
        if not expected_months or reported.metric_value is None:
            continue

        observations = [
            fact
            for fact in facts
            if fact.metric_name == reported.metric_name
            and fact.unit == reported.unit
            and fact.aggregation_role == "observation"
            and fact.period_grain == "month"
            and fact.period_start in expected_months
            and fact.metric_value is not None
        ]
        facts_by_month: dict[date, list[FactCandidate]] = {}
        for fact in observations:
            assert fact.period_start is not None
            facts_by_month.setdefault(fact.period_start, []).append(fact)
        if set(facts_by_month) != set(expected_months):
            continue

        selected_month_facts: list[FactCandidate] = []
        for month_start in expected_months:
            candidates = facts_by_month[month_start]
            if len({fact.metric_value for fact in candidates}) > 1:
                warnings.append(
                    f"Conflicting monthly {reported.metric_name} values for "
                    f"{month_start:%B %Y} appear in multiple workbook sheets."
                )
            candidates.sort(key=_monthly_source_priority)
            selected_month_facts.append(candidates[0])

        report_sheet = reported.evidence.sheet_name or "workbook"
        period_key = (
            reported.metric_name,
            reported.unit,
            reported.period_start,
            reported.period_end,
        )
        calculated_fact = calculated_by_period.get(period_key)
        if calculated_fact is None:
            calculated_total = sum(
                (fact.metric_value or Decimal(0) for fact in selected_month_facts),
                Decimal(0),
            )
            calculated_fact = FactCandidate(
                metric_name=reported.metric_name,
                fact_kind="system_calculated",
                aggregation_role="system_total",
                raw_value=str(calculated_total),
                metric_value=calculated_total,
                unit=reported.unit,
                currency_code=reported.currency_code,
                period_start=reported.period_start,
                period_end=reported.period_end,
                period_grain=reported.period_grain,
                financial_year=reported.financial_year,
                financial_year_start_year=reported.financial_year_start_year,
                dimensions={"calculation_basis": "monthly_source_facts"},
                confidence=Decimal("1.000"),
                value_origin="system_calculated",
                calculation_method=(
                    f"sum of {len(selected_month_facts)} monthly observations"
                ),
                derived_from=[fact.id for fact in selected_month_facts],
                source_reference=(
                    f"Calculated from monthly values for {reported.period_start} "
                    f"to {reported.period_end}"
                ),
                evidence=EvidenceCandidate(
                    evidence_kind="calculation",
                    quote=(
                        f"Sum of {len(selected_month_facts)} monthly "
                        f"{reported.metric_name} values: {calculated_total} "
                        f"{reported.unit or ''}."
                    ),
                    source_metadata={
                        "operation": "sum_monthly_values",
                        "month_count": len(selected_month_facts),
                    },
                ),
            )
            calculated_by_period[period_key] = calculated_fact
            additions.append(calculated_fact)

        delta = (calculated_fact.metric_value or Decimal(0)) - reported.metric_value
        delta_description = (
            "Monthly sum matches the reported total."
            if delta == 0
            else f"Monthly sum differs from the reported total by {delta} {reported.unit or ''}."
        )
        additions.append(
            FactCandidate(
                metric_name="reconciliation_delta",
                fact_kind="reconciliation",
                aggregation_role="reconciliation_delta",
                raw_value=str(delta),
                metric_value=delta,
                unit=reported.unit,
                currency_code=reported.currency_code,
                period_start=reported.period_start,
                period_end=reported.period_end,
                period_grain=reported.period_grain,
                dimensions={
                    "metric_name": reported.metric_name,
                    "sheet_name": report_sheet,
                    "reported_fact_id": str(reported.id),
                },
                confidence=Decimal("1.000"),
                value_origin="system_calculated",
                calculation_method="monthly sum minus source-reported total",
                derived_from=[calculated_fact.id, reported.id],
                description=delta_description,
                source_reference=(
                    f"Reconciliation for {reported.metric_name} in {report_sheet}"
                ),
                evidence=EvidenceCandidate(
                    evidence_kind="calculation",
                    quote=(
                        f"{calculated_fact.metric_value} monthly sum - "
                        f"{reported.metric_value} reported total = {delta}."
                    ),
                    source_metadata={
                        "operation": "calculated_minus_reported",
                        "reported_fact_id": str(reported.id),
                        "calculated_fact_id": str(calculated_fact.id),
                        "reported_sheet": report_sheet,
                    },
                ),
            )
        )
        if delta != 0:
            warnings.append(
                f"{report_sheet}: reported {reported.metric_name} differs from "
                f"the monthly sum by {delta} {reported.unit or ''} for "
                f"{reported.period_start} to {reported.period_end}."
            )

    facts.extend(additions)
    return list(dict.fromkeys(warnings))


def _months_in_period(period_start: date, period_end: date) -> list[date]:
    months: list[date] = []
    year, month = period_start.year, period_start.month
    while (year, month) <= (period_end.year, period_end.month):
        month_start = date(year, month, 1)
        month_end = date(year, month, monthrange(year, month)[1])
        if month_start >= period_start and month_end <= period_end:
            months.append(month_start)
        month += 1
        if month == 13:
            year += 1
            month = 1
    return months


def _monthly_source_priority(fact: FactCandidate) -> tuple[int, str]:
    sheet_name = (fact.evidence.sheet_name or "").casefold()
    if "monthly summary" in sheet_name:
        priority = 0
    elif "monthly trends" in sheet_name:
        priority = 1
    else:
        priority = 2
    return priority, fact.source_reference


def _financial_year_coverage_warnings(facts: list[FactCandidate]) -> list[str]:
    months_by_year: dict[int, set[date]] = {}
    for fact in facts:
        if (
            fact.aggregation_role != "observation"
            or fact.period_grain != "month"
            or fact.period_start is None
            or fact.financial_year_start_year is None
        ):
            continue
        months_by_year.setdefault(fact.financial_year_start_year, set()).add(
            fact.period_start
        )

    warnings: list[str] = []
    for year, month_starts in sorted(months_by_year.items()):
        if len(month_starts) == 12:
            continue
        month_names = ", ".join(
            month_start.strftime("%b")
            for month_start in sorted(month_starts)
        )
        warnings.append(
            f"{financial_year_label(year)} has monthly observations for "
            f"{len(month_starts)}/12 months ({month_names}); FY coverage is partial."
        )
    return warnings


def _workbook_year_context(workbook) -> int | None:
    title_years: set[int] = set()
    contextual_years: set[int] = set()
    for worksheet in workbook.worksheets:
        title_years.update(int(match) for match in _YEAR_PATTERN.findall(worksheet.title))
        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                match = _YEAR_CONTEXT_PATTERN.search(cell.value)
                if match:
                    contextual_years.add(int(match.group(1)))
    if len(title_years) == 1:
        return next(iter(title_years))
    if len(contextual_years) == 1:
        return next(iter(contextual_years))
    return None


def _sheet_year(worksheet, workbook_year: int | None) -> int | None:
    years = {int(match) for match in _YEAR_PATTERN.findall(worksheet.title)}
    if len(years) == 1:
        return next(iter(years))
    return workbook_year


def _header_unit(header: str) -> tuple[str | None, str | None]:
    normalized = _normalize_header(header)
    if "inr" in normalized or "rupee" in normalized or "rs " in normalized:
        if "per kwh" in normalized or "/kwh" in normalized:
            return "INR/kWh", "INR"
        return "INR", "INR"
    if "%" in header or "percent" in normalized or "share" in normalized:
        return "%", None
    if "kwh" in normalized:
        return "kWh", None
    if "mwh" in normalized:
        return "MWh", None
    if "wh" in normalized:
        return "Wh", None
    if "kva" in normalized:
        return "kVA", None
    if "kwp" in normalized:
        return "kWp", None
    if re.search(r"\bkg\b", normalized):
        return "kg", None
    if "tonne" in normalized or "tons" in normalized or re.search(r"\bmt\b", normalized):
        return "tonnes", None
    if "cubic meter" in normalized or "cu m" in normalized:
        return "m³", None
    if "liter" in normalized or "litre" in normalized:
        return "liters", None
    if "unit rate" in normalized:
        return None, None
    return None, None


def _spreadsheet_decimal(value: Any) -> tuple[Decimal | None, str | None]:
    if value is None or isinstance(value, (bool, date, datetime)):
        return None, None
    unit: str | None = None
    if isinstance(value, Decimal):
        parsed = value
    elif isinstance(value, (int, float)):
        parsed = Decimal(str(value))
    else:
        text = str(value).strip()
        if not text or text in {"-", "—", "–", "N/A", "n/a"}:
            return None, None
        if "%" in text:
            unit = "%"
        parsed = _parse_decimal(text)
    if parsed is None or not parsed.is_finite():
        return None, unit
    _validate_decimal_precision(parsed)
    return parsed, unit


def _parse_decimal(value: str | int | float | Decimal) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    if not text:
        return None
    negative = (text.startswith("(") and text.endswith(")")) or text.casefold().endswith("cr")
    cleaned = re.sub(r"(?i)\b(cr|inr|rs\.?|kwh|mwh|wh|kva|kvah|kg|mt|tonnes?|tons?|cubic\s*meters?|cu\.?m)\b", "", text)
    cleaned = cleaned.replace("₹", "").replace("$", "")
    cleaned = cleaned.replace(",", "").replace("%", "").strip()
    cleaned = cleaned.strip("() ")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", cleaned)
    if not match:
        return None
    try:
        parsed = Decimal(match.group(0))
    except InvalidOperation:
        return None
    parsed = -abs(parsed) if negative else parsed
    _validate_decimal_precision(parsed)
    return parsed


def _validate_decimal_precision(value: Decimal) -> None:
    normalized = value.normalize()
    parts = normalized.as_tuple()
    scale = max(0, -parts.exponent)
    integer_digits = max(0, len(parts.digits) + parts.exponent)
    if scale > 12 or integer_digits > 26:
        raise ValueError("Extracted value exceeds supported numeric precision.")


def _parse_date(value: str | date | datetime | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = re.sub(r"(?<=\d)(?:st|nd|rd|th)\b", "", str(value).strip(), flags=re.I)
    for date_format in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d.%m.%Y",
        "%m/%d/%Y",
        "%Y/%m/%d",
        "%d %B %Y",
        "%d %b %Y",
        "%B %d, %Y",
        "%b %d, %Y",
    ):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    for date_format in ("%B %Y", "%b %Y", "%Y-%m"):
        try:
            parsed = datetime.strptime(text, date_format).date()
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            continue
    return None


def _month_number(value: str) -> int | None:
    return _MONTHS.get(value.strip().casefold())


def _period_grain(start: date | None, end: date | None) -> str | None:
    if start is None or end is None:
        return None
    if start.day == 1 and end == date(start.year, start.month, monthrange(start.year, start.month)[1]):
        return "month"
    if start.month == 1 and start.day == 1 and end.month == 12 and end.day == 31:
        return "year"
    return "range"


def _first_text(values: list[Any]) -> str | None:
    for value in values:
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def _cell_to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _is_empty_marker(value: Any) -> bool:
    return isinstance(value, str) and value.strip() in {"", "-", "—", "–", "N/A", "n/a"}


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9%]+", " ", value.casefold()).strip()


def _slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")[:100]


def _excel_sheet_reference(sheet_name: str) -> str:
    escaped = sheet_name.replace("'", "''")
    return f"'{escaped}'"


def _quote_is_present(quote: str, source_text: str) -> bool | None:
    if not source_text.strip():
        return None
    normalize = lambda text: re.sub(r"\s+", " ", text).strip().casefold()
    return normalize(quote) in normalize(source_text)
