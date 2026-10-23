import os
import unittest
from contextlib import nullcontext
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

from openpyxl import Workbook
from pypdf import PdfWriter

from schemas.documents import DocumentResponse, sanitize_document_error
from services.fact_extraction import (
    EvidenceCandidate,
    ExtractedDocument,
    FactCandidate,
    extract_pdf_facts,
    extract_spreadsheet_facts,
    _parse_date,
    _parse_decimal,
)
from services.fact_persistence import persist_extracted_facts
from services.llm_observability import langfuse_observation


class FactExtractionTests(unittest.TestCase):
    def test_monthly_workbook_preserves_cell_sources_and_fy(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Monthly Summary 2026"
        sheet.append(["Annual Electricity Summary"])
        sheet.append(
            [
                "Month",
                "Grid Supply (kWh)",
                "Solar Power (kWh)",
                "Wind Power (kWh)",
                "Total Renewable (kWh)",
                "Total Consumption (kWh)",
                "Renewable %",
            ]
        )
        months = (
            "January",
            "February",
            "March",
            "April",
            "May",
            "June",
            "July",
            "August",
            "September",
            "October",
            "November",
            "December",
        )
        for month in months:
            sheet.append([month, 10, 2, 1, 3, 13, "23.1%"])
        sheet.append([None])
        sheet.append(["ANNUAL TOTAL", 120, 24, 12, 36, 155, "23.1%"])

        buffer = BytesIO()
        workbook.save(buffer)
        extracted = extract_spreadsheet_facts(buffer.getvalue())

        january_grid = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "grid_electricity_consumption"
            and fact.period_start == date(2026, 1, 1)
        )
        april_grid = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "grid_electricity_consumption"
            and fact.period_start == date(2026, 4, 1)
        )
        renewable = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "renewable_generation"
            and fact.period_start == date(2026, 1, 1)
        )
        wind = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "wind_generation"
            and fact.period_start == date(2026, 1, 1)
        )
        monthly_percent = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "renewable_share_percent"
            and fact.period_start == date(2026, 1, 1)
        )
        annual_total = next(
            fact
            for fact in extracted.facts
            if fact.metric_name == "total_energy_consumption"
            and fact.aggregation_role == "reported_total"
        )
        delta = next(
            fact
            for fact in extracted.facts
            if fact.fact_kind == "reconciliation"
            and fact.dimensions.get("metric_name") == "total_energy_consumption"
        )

        self.assertEqual(january_grid.source_reference, "'Monthly Summary 2026'!B3")
        self.assertEqual(january_grid.financial_year, "FY 2025-26")
        self.assertEqual(april_grid.financial_year, "FY 2026-27")
        self.assertEqual(renewable.metric_name, "renewable_generation")
        self.assertEqual(wind.metric_name, "wind_generation")
        self.assertEqual(monthly_percent.metric_value, Decimal("23.1"))
        self.assertEqual(monthly_percent.unit, "%")
        self.assertEqual(annual_total.financial_year_start_year, None)
        self.assertEqual(delta.metric_value, 1)
        self.assertTrue(any("FY 2025-26" in warning for warning in extracted.warnings))

    def test_uncached_formula_is_not_misread_as_a_number(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Monthly Summary 2026"
        sheet.append(["Month", "Total Consumption (kWh)"])
        sheet.append(["January", "=SUM(100,200)"])

        buffer = BytesIO()
        workbook.save(buffer)
        extracted = extract_spreadsheet_facts(buffer.getvalue())
        formula_fact = next(
            fact for fact in extracted.facts if fact.period_start == date(2026, 1, 1)
        )

        self.assertEqual(formula_fact.fact_kind, "formula")
        self.assertIsNone(formula_fact.metric_value)
        self.assertEqual(formula_fact.raw_value, "=SUM(100,200)")

    def test_pdf_facts_keep_page_quote_and_raw_value(self) -> None:
        class FakeStructuredModel:
            async def ainvoke(self, _messages):
                return {
                    "document_kind": "utility_bill",
                    "entity_name": "Example Factory",
                    "facts": [
                        {
                            "metric_name": "previous_meter_reading",
                            "fact_kind": "meter_reading",
                            "aggregation_role": "observation",
                            "value_type": "numeric",
                            "raw_value": "12,500",
                            "unit": "Cubic Meters",
                            "currency_code": None,
                            "event_date": "01-12-2025",
                            "period_start": None,
                            "period_end": None,
                            "period_grain": None,
                            "description": "Previous meter reading",
                            "dimensions": {"meter_id": "MTR-1", "reading_type": "previous"},
                            "page_number": 1,
                            "source_quote": "Previous Reading (01-12-2025): 12,500 Cubic Meters",
                            "confidence": 0.99,
                        }
                    ],
                }

        class FakeChatModel:
            def __init__(self, **_kwargs):
                pass

            def with_structured_output(self, _schema, **_kwargs):
                return FakeStructuredModel()

        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        pdf_buffer = BytesIO()
        writer.write(pdf_buffer)

        async def run_extraction():
            return await extract_pdf_facts(pdf_buffer.getvalue())

        with (
            patch.dict(os.environ, {"GEMINI_API_KEY": "test-key"}),
            patch("langchain_google_genai.ChatGoogleGenerativeAI", FakeChatModel),
            patch(
                "services.fact_extraction.langfuse_observation",
                side_effect=lambda **_kwargs: nullcontext(None),
            ),
        ):
            import asyncio

            extracted = asyncio.run(run_extraction())

        fact = extracted.facts[0]
        self.assertEqual(extracted.document_kind, "utility_bill")
        self.assertEqual(extracted.entity_name, "Example Factory")
        self.assertEqual(fact.metric_value, Decimal("12500"))
        self.assertEqual(fact.raw_value, "12,500")
        self.assertEqual(fact.source_reference, "PDF page 1")
        self.assertEqual(fact.evidence.page_number, 1)
        self.assertIn("Previous Reading", fact.evidence.quote)

    def test_decimal_and_date_normalization_keep_source_values_separate(self) -> None:
        self.assertEqual(_parse_decimal("₹1,23,456.75"), 123456.75)
        self.assertEqual(_parse_decimal("500.00Cr"), -500)
        self.assertEqual(_parse_date("31st August 2026"), date(2026, 8, 31))
        self.assertEqual(_parse_date("10.03.2023"), date(2023, 3, 10))

    def test_internal_database_errors_are_not_user_facing(self) -> None:
        internal = (
            "sqlalchemy.exc.IntegrityError: violates foreign key constraint "
            "[SQL: INSERT INTO extracted_records]"
        )
        self.assertEqual(
            sanitize_document_error(internal),
            "Document processing failed. Please try again or contact support.",
        )
        self.assertEqual(
            sanitize_document_error("This PDF is password-protected and cannot be processed."),
            "This PDF is password-protected and cannot be processed.",
        )
        response = DocumentResponse(
            id=uuid4(),
            user_id=1,
            filename="bill.pdf",
            cloudinary_url=None,
            cloudinary_public_id="documents/1/bill.pdf",
            status="failed",
            error_message=internal,
            upload_timestamp=datetime.now(UTC),
        )
        self.assertEqual(
            response.error_message,
            "Document processing failed. Please try again or contact support.",
        )

    def test_source_evidence_is_flushed_before_fact_foreign_keys(self) -> None:
        class RecordingSession:
            def __init__(self):
                self.events = []
                self.rows = []

            def add_all(self, rows):
                self.rows.extend(rows)
                self.events.append(("add_all", {type(row).__name__ for row in rows}))

            async def flush(self):
                self.events.append(("flush", set()))

        session = RecordingSession()
        extracted = ExtractedDocument(
            document_kind="utility_bill",
            facts=[
                FactCandidate(
                    metric_name="reported_consumption",
                    metric_value=Decimal("2321"),
                    raw_value="2\x00321",
                    unit="m³",
                    dimensions={"meter_id": "MTR\x00-1"},
                    source_reference="PDF page 1",
                    evidence=EvidenceCandidate(
                        evidence_kind="pdf_page",
                        page_number=1,
                        quote="Units\x00 consumed: 2,321 cubic meters",
                    ),
                )
            ],
        )

        async def persist():
            await persist_extracted_facts(
                session,
                document_id=uuid4(),
                run=type("Run", (), {"id": uuid4()})(),
                document_kind="utility_bill",
                extracted=extracted,
            )

        import asyncio

        asyncio.run(persist())
        self.assertEqual(session.events[0][1], {"SourceEvidence"})
        self.assertEqual(session.events[1][0], "flush")
        self.assertEqual(session.events[2][1], {"ExtractedRecord"})
        evidence = next(row for row in session.rows if type(row).__name__ == "SourceEvidence")
        fact = next(row for row in session.rows if type(row).__name__ == "ExtractedRecord")
        self.assertNotIn("\x00", evidence.quote)
        self.assertNotIn("\x00", fact.raw_value)
        self.assertNotIn("\x00", fact.dimensions["meter_id"])

    def test_langfuse_observation_wraps_a_generation_when_configured(self) -> None:
        class FakeObservation:
            def __init__(self):
                self.updates = []

            def update(self, **values):
                self.updates.append(values)

        class FakeContext:
            def __init__(self, observation):
                self.observation = observation

            def __enter__(self):
                return self.observation

            def __exit__(self, *_args):
                return False

        class FakeClient:
            def __init__(self):
                self.attributes = None
                self.observation = FakeObservation()

            def start_as_current_observation(self, **attributes):
                self.attributes = attributes
                return FakeContext(self.observation)

        client = FakeClient()
        with patch(
            "services.llm_observability.get_langfuse_client", return_value=client
        ):
            with langfuse_observation(
                as_type="generation",
                name="test-generation",
                model="gemini-test",
            ) as observation:
                observation.update(output={"fact_count": 1})

        self.assertEqual(client.attributes["name"], "test-generation")
        self.assertEqual(
            client.observation.updates,
            [{"output": {"fact_count": 1}}],
        )


if __name__ == "__main__":
    unittest.main()
